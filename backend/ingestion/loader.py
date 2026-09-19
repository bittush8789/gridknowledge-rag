import os
import re
import hashlib
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

Tuple_Dict_Text = Tuple[Dict[str, str], str]

# Optional imports for rich document parsing
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None


class LoadedPage:
    def __init__(self, page_number: int, section: str, text: str):
        self.page_number = page_number
        self.section = section
        self.text = text

    def to_dict(self) -> Dict[str, Any]:
        return {
            "page_number": self.page_number,
            "section": self.section,
            "text": self.text,
        }


class LoadedDocument:
    def __init__(
        self,
        title: str,
        filename: str,
        category: str,
        equipment: Optional[str],
        section: Optional[str],
        version: str,
        effective_date: Optional[str],
        status: str,
        access_level: str,
        file_path: str,
        file_hash: str,
        pages: List[LoadedPage],
    ):
        self.title = title
        self.filename = filename
        self.category = category
        self.equipment = equipment
        self.section = section
        self.version = version
        self.effective_date = effective_date
        self.status = status
        self.access_level = access_level
        self.file_path = file_path
        self.file_hash = file_hash
        self.pages = pages


class DocumentLoader:
    """Multi-format enterprise document loader with page and section preservation."""

    def compute_file_hash(self, file_path: Path) -> str:
        hasher = hashlib.sha256()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                hasher.update(chunk)
        return hasher.hexdigest()

    def parse_frontmatter(self, text: str) -> Tuple_Dict_Text:
        """Extract YAML/header frontmatter if present."""
        metadata = {}
        content = text

        match = re.match(r"^---\s*\n(.*?)\n---\s*\n(.*)$", text, re.DOTALL)
        if match:
            fm_text = match.group(1)
            content = match.group(2)
            for line in fm_text.splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    k = k.strip().lower()
                    v = v.strip().strip('"').strip("'")
                    metadata[k] = v
        return metadata, content

    def load_markdown_or_text(self, file_path: Path) -> LoadedDocument:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            raw_text = f.read()

        metadata, content = self.parse_frontmatter(raw_text)

        # Derive defaults
        filename = file_path.name
        title = metadata.get("document") or metadata.get("title") or file_path.stem.replace("_", " ").title()
        category = metadata.get("category") or file_path.parent.name.replace("_", " ").title()
        equipment = metadata.get("equipment", "General Grid Asset")
        version = metadata.get("version", "1.0")
        effective_date = metadata.get("effective_date", "2024-01-01")
        status = metadata.get("status", "active")
        access_level = metadata.get("access_level", "Viewer")

        # Parse pages by markers like "--- Page X ---"
        page_splits = re.split(r"(?:^|\n)---\s*Page\s*(\d+)\s*---", content)
        pages: List[LoadedPage] = []

        if len(page_splits) > 1:
            # page_splits format: [before_page_1, page_num_1, text_1, page_num_2, text_2, ...]
            initial_text = page_splits[0].strip()
            if initial_text:
                pages.append(LoadedPage(page_number=1, section="Overview", text=initial_text))

            for i in range(1, len(page_splits), 2):
                p_num = int(page_splits[i])
                p_text = page_splits[i + 1].strip()

                # Extract section name if present
                sec_match = re.search(r"##\s*Section\s*\d*[:\s]*([^\n]+)", p_text)
                sec_name = sec_match.group(1).strip() if sec_match else f"Page {p_num}"

                pages.append(LoadedPage(page_number=p_num, section=sec_name, text=p_text))
        else:
            # Single page document
            pages.append(LoadedPage(page_number=1, section=title, text=content.strip()))

        file_hash = self.compute_file_hash(file_path)

        return LoadedDocument(
            title=title,
            filename=filename,
            category=category,
            equipment=equipment,
            section=pages[0].section if pages else "General",
            version=version,
            effective_date=effective_date,
            status=status,
            access_level=access_level,
            file_path=str(file_path.resolve()),
            file_hash=file_hash,
            pages=pages,
        )

    def load_pdf(self, file_path: Path) -> LoadedDocument:
        if not fitz:
            raise ImportError("PyMuPDF (fitz) is required for PDF parsing.")

        doc = fitz.open(file_path)
        pages: List[LoadedPage] = []

        for p_idx, page in enumerate(doc):
            p_text = page.get_text("text").strip()
            p_num = p_idx + 1
            sec_match = re.search(r"(?:Section|Chapter)\s*\d*[:\s]*([^\n]+)", p_text, re.IGNORECASE)
            sec_name = sec_match.group(1).strip() if sec_match else f"Page {p_num}"
            pages.append(LoadedPage(page_number=p_num, section=sec_name, text=p_text))

        doc.close()
        file_hash = self.compute_file_hash(file_path)
        title = file_path.stem.replace("_", " ").title()

        return LoadedDocument(
            title=title,
            filename=file_path.name,
            category=file_path.parent.name.replace("_", " ").title(),
            equipment="General",
            section="General",
            version="1.0",
            effective_date="2024-01-01",
            status="active",
            access_level="Viewer",
            file_path=str(file_path.resolve()),
            file_hash=file_hash,
            pages=pages,
        )

    def load_docx(self, file_path: Path) -> LoadedDocument:
        if not docx:
            raise ImportError("python-docx is required for DOCX parsing.")

        doc = docx.Document(file_path)
        text_lines = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(text_lines)

        pages = [LoadedPage(page_number=1, section="General", text=full_text)]
        file_hash = self.compute_file_hash(file_path)

        return LoadedDocument(
            title=file_path.stem.replace("_", " ").title(),
            filename=file_path.name,
            category=file_path.parent.name.replace("_", " ").title(),
            equipment="General",
            section="General",
            version="1.0",
            effective_date="2024-01-01",
            status="active",
            access_level="Viewer",
            file_path=str(file_path.resolve()),
            file_hash=file_hash,
            pages=pages,
        )

    def load_html(self, file_path: Path) -> LoadedDocument:
        if not BeautifulSoup:
            raise ImportError("BeautifulSoup4 is required for HTML parsing.")

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            soup = BeautifulSoup(f.read(), "html.parser")

        title_tag = soup.find("title")
        title = title_tag.get_text().strip() if title_tag else file_path.stem.replace("_", " ").title()
        body_text = soup.get_text(separator="\n\n").strip()

        pages = [LoadedPage(page_number=1, section=title, text=body_text)]
        file_hash = self.compute_file_hash(file_path)

        return LoadedDocument(
            title=title,
            filename=file_path.name,
            category=file_path.parent.name.replace("_", " ").title(),
            equipment="General",
            section="General",
            version="1.0",
            effective_date="2024-01-01",
            status="active",
            access_level="Viewer",
            file_path=str(file_path.resolve()),
            file_hash=file_hash,
            pages=pages,
        )

    def load_file(self, file_path: Path) -> LoadedDocument:
        suffix = file_path.suffix.lower()
        if suffix in [".md", ".txt"]:
            return self.load_markdown_or_text(file_path)
        elif suffix == ".pdf":
            return self.load_pdf(file_path)
        elif suffix in [".docx", ".doc"]:
            return self.load_docx(file_path)
        elif suffix in [".html", ".htm"]:
            return self.load_html(file_path)
        else:
            return self.load_markdown_or_text(file_path)

    def scan_directory(self, base_dir: Path) -> List[LoadedDocument]:
        """Recursively scans data/documents/ and loads all supported files."""
        loaded_docs: List[LoadedDocument] = []
        for root, _, files in os.walk(base_dir):
            for file in files:
                file_path = Path(root) / file
                if file_path.suffix.lower() in [".md", ".txt", ".pdf", ".docx", ".html"]:
                    try:
                        doc = self.load_file(file_path)
                        loaded_docs.append(doc)
                    except Exception as e:
                        print(f"[DocumentLoader] Error loading {file_path}: {e}")
        return loaded_docs


loader = DocumentLoader()
